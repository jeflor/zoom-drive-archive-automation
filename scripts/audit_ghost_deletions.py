import os
import sys, io, json
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from collections import Counter
import openpyxl

ROOT = os.environ['DRIVE_ROOT_FOLDER_ID']

zoom = json.load(open('zoom_recordings.json'))
print(f"Live-scan meetings: {len(zoom)}")

# 1) duplicate UUIDs?
uuids = [m['uuid'] for m in zoom]
dups = [u for u,c in Counter(uuids).items() if c>1]
print(f"Duplicate UUIDs in my scan: {len(dups)}")

# 2) total_size vs sum of files
t_field = sum(m['total_size'] for m in zoom)
t_files = sum(sum(f['size'] for f in m['files']) for m in zoom)
print(f"Sum of total_size field: {t_field/1e9:.1f} GB")
print(f"Sum of per-file sizes  : {t_files/1e9:.1f} GB")

# 3) pull DELETED uuids from the final-status spreadsheet
creds = Credentials.from_authorized_user_file(os.environ.get('DRIVE_TOKEN_JSON', 'token.json'), ['https://www.googleapis.com/auth/drive'])
if not creds.valid: creds.refresh(Request())
svc = build('drive','v3',credentials=creds,cache_discovery=False)
r=svc.files().list(q=f"'{ROOT}' in parents and name='prior_audit.xlsx' and trashed=false",
    fields='files(id)', supportsAllDrives=True, includeItemsFromAllDrives=True, corpora='allDrives').execute()
buf=io.BytesIO(); dl=MediaIoBaseDownload(buf, svc.files().get_media(fileId=r['files'][0]['id'], supportsAllDrives=True))
done=False
while not done: _,done=dl.next_chunk()
buf.seek(0); wb=openpyxl.load_workbook(buf, read_only=True, data_only=True)
ws=wb['Deletion Candidates (2TB)']
hdr=list(next(ws.iter_rows(max_row=1, values_only=True)))
iU=hdr.index('Meeting UUID'); iD=hdr.index('Deletion Status'); iSz=hdr.index('Size (GB)')
deleted_uuids=set(); deleted_gb=0
for row in ws.iter_rows(min_row=2, values_only=True):
    if row and row[iU] and 'ALREADY DELETED' in str(row[iD]):
        deleted_uuids.add(row[iU]); deleted_gb += row[iSz] or 0
print(f"\nSpreadsheet says ALREADY DELETED: {len(deleted_uuids)} recordings, {deleted_gb:.0f} GB")

live=set(uuids)
still_live = deleted_uuids & live
print(f"Of those, STILL appearing in my live scan: {len(still_live)}")
gb_ghost = sum((m['total_size'] or 0) for m in zoom if m['uuid'] in deleted_uuids)/1e9
print(f"GB they contribute to my 2034 total: {gb_ghost:.0f} GB")
print(f"\nMy total minus these trashed ghosts: {(t_field/1e9 - gb_ghost):.0f} GB  (Zoom bill says ~1240 GB)")
