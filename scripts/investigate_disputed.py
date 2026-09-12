import os
import sys, io, json
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from collections import defaultdict, Counter
import openpyxl

ROOT = os.environ['DRIVE_ROOT_FOLDER_ID']

zoom = {m['uuid']: m for m in json.load(open('zoom_recordings.json'))}
drive = json.load(open('drive_index.json'))
folder_files = defaultdict(list)
for f in drive['files']:
    folder_files[f['path']].append(f)

wb_mine = openpyxl.load_workbook('safe_to_delete.xlsx', read_only=True, data_only=True)
ws = wb_mine['Confirmed Safe (media)']
hdr = list(next(ws.iter_rows(max_row=1, values_only=True)))
iu, ip = hdr.index('uuid'), hdr.index('drive_full_path')
mine = {row[iu]: row[ip] for row in ws.iter_rows(min_row=2, values_only=True) if row and row[iu]}

# previous verification detail
creds = Credentials.from_authorized_user_file(os.environ.get('DRIVE_TOKEN_JSON', 'token.json'), ['https://www.googleapis.com/auth/drive'])
if not creds.valid: creds.refresh(Request())
svc = build('drive','v3',credentials=creds,cache_discovery=False)
r=svc.files().list(q=f"'{ROOT}' in parents and name='prior_audit.xlsx' and trashed=false",
    fields='files(id)', supportsAllDrives=True, includeItemsFromAllDrives=True, corpora='allDrives').execute()
buf=io.BytesIO(); dl=MediaIoBaseDownload(buf, svc.files().get_media(fileId=r['files'][0]['id'], supportsAllDrives=True))
done=False
while not done: _,done=dl.next_chunk()
buf.seek(0); wb=openpyxl.load_workbook(buf, read_only=True, data_only=True)
ws2=wb['Deletion Candidates (2TB)']
h2=list(next(ws2.iter_rows(max_row=1, values_only=True)))
iU,iDet,iV=h2.index('Meeting UUID'),h2.index('Verification Detail'),h2.index('Verification Status')
prev={row[iU]:(row[iV],row[iDet]) for row in ws2.iter_rows(min_row=2, values_only=True) if row and row[iU]}

overlap=[u for u in mine if u in prev]
print(f"Investigating {len(overlap)} disputed meetings\n"+"="*90)
for u in overlap:
    m=zoom[u]; path=mine[u]
    cloud_media=[f for f in m['files'] if f['file_type'] in ('MP4','M4A') and f['size']>=1_000_000]
    dfiles=folder_files.get(path,[])
    dsizes=Counter(f['size'] for f in dfiles)
    print(f"\n{m['start_time'][:10]}  {m['topic'][:55]}")
    print(f"  prev verdict: {prev[u][0]} — {str(prev[u][1])[:95]}")
    print(f"  Drive folder: {path}")
    print(f"  Drive folder has {len(dfiles)} files")
    avail=Counter(dsizes)
    for f in sorted(cloud_media,key=lambda x:-x['size']):
        hit = avail[f['size']]>0
        if hit: avail[f['size']]-=1
        print(f"    CLOUD {f['recording_type']:38} {f['size']:>13,}  {'MATCH' if hit else 'NO MATCH in folder'}")
    # show drive folder media files
    print(f"    -- Drive folder media files --")
    for f in sorted([x for x in dfiles if x['size']>=1_000_000], key=lambda x:-x['size'])[:10]:
        print(f"    DRIVE {f['name'][:50]:50} {f['size']:>13,}")
