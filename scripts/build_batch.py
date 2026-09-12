import os
import sys, io, json
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from collections import defaultdict, Counter
import openpyxl

ROOT = os.environ['DRIVE_ROOT_FOLDER_ID']

zoom = json.load(open('zoom_recordings.json'))
drive = json.load(open('drive_index.json'))
MEDIA={'MP4','M4A'}
folder_sizes=defaultdict(Counter); size_paths=defaultdict(list); gsize=Counter()
for f in drive['files']:
    folder_sizes[f['path']][f['size']]+=1; size_paths[f['size']].append(f['path']); gsize[f['size']]+=1
def best_folder(media):
    sizes=sorted((f['size'] for f in media),reverse=True); cand=set(size_paths.get(sizes[0],[]))
    best,bn=None,-1
    for p in cand:
        av=Counter(folder_sizes[p]); n=0
        for s in sizes:
            if av[s]>0: av[s]-=1; n+=1
        if n>bn: best,bn=p,n
    return best,bn

confirmed={}
for m in zoom:
    media=[f for f in m['files'] if f['file_type'] in MEDIA and f['size']>=1_000_000]
    if not media: continue
    need=Counter(f['size'] for f in media)
    if any(gsize[s]<need[s] for s in need): continue
    bf,n=best_folder(media)
    if n!=len(media): continue
    confirmed[m['uuid']]={'uuid':m['uuid'],'host':m['host'],'date':(m['start_time'] or '')[:10],
        'topic':m['topic'],'drive_folder':bf,
        'files':[{'id':f['id'],'recording_type':f['recording_type'],'size':f['size']} for f in media]}

# 927 sheet candidate UUIDs + prev detail
creds=Credentials.from_authorized_user_file(os.environ.get('DRIVE_TOKEN_JSON', 'token.json'),['https://www.googleapis.com/auth/drive'])
if not creds.valid: creds.refresh(Request())
svc=build('drive','v3',credentials=creds,cache_discovery=False)
r=svc.files().list(q=f"'{ROOT}' in parents and name='prior_audit.xlsx' and trashed=false",
    fields='files(id)',supportsAllDrives=True,includeItemsFromAllDrives=True,corpora='allDrives').execute()
buf=io.BytesIO(); dl=MediaIoBaseDownload(buf, svc.files().get_media(fileId=r['files'][0]['id'],supportsAllDrives=True))
d=False
while not d: _,d=dl.next_chunk()
buf.seek(0); wb=openpyxl.load_workbook(buf,read_only=True,data_only=True); ws=wb['Deletion Candidates (2TB)']
h=list(next(ws.iter_rows(max_row=1,values_only=True)))
iU,iV,iDet=h.index('Meeting UUID'),h.index('Verification Status'),h.index('Verification Detail')
sheet={row[iU]:(row[iV],row[iDet]) for row in ws.iter_rows(min_row=2,values_only=True) if row and row[iU]}

disputed=[u for u in confirmed if u in sheet]
batch=[confirmed[u] for u in confirmed if u not in disputed]
json.dump(batch, open('batch_delete.json','w'))

# hold-list spreadsheet for the 11
hb=openpyxl.Workbook(); hs=hb.active; hs.title='Hold - confirm manually'
hs.append(['host','date','topic','media_files','media_GB','drive_folder','prev_verdict','prev_detail','uuid'])
for u in disputed:
    c=confirmed[u]; gb=round(sum(f['size'] for f in c['files'])/1e9,2)
    hs.append([c['host'],c['date'],c['topic'][:80],len(c['files']),gb,c['drive_folder'],
               sheet[u][0],str(sheet[u][1])[:120],u])
for col,w in {'A':22,'B':12,'C':46,'D':11,'E':10,'F':64,'G':10,'H':70,'I':30}.items(): hs.column_dimensions[col].width=w
hb.save('disputed_hold_list.xlsx')

bfiles=sum(len(c['files']) for c in batch); bgb=sum(sum(f['size'] for f in c['files']) for c in batch)/1e9
print(f"BATCH to trash: {len(batch)} meetings, {bfiles} media files, {bgb:.1f} GB")
print(f"HELD for manual confirm: {len(disputed)} meetings -> disputed_hold_list.xlsx")
