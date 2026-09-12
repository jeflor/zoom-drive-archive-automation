"""Rebuild duplicates.xlsx to show, per duplicate group, which copy was KEPT
and which were TRASHED. Reflects actual result from drive_trashed_log.txt."""
import sys, json
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill

files=json.load(open('drive_index_md5.json'))
trashed=set(l.strip() for l in open('drive_trashed_log.txt')) if __import__('os').path.exists('drive_trashed_log.txt') else set()
by_md5=defaultdict(list)
for f in files:
    if f['md5']: by_md5[f['md5']].append(f)
groups=[(m,fs) for m,fs in by_md5.items() if len(fs)>1]
groups.sort(key=lambda g:-(g[1][0]['size']*(len(g[1])-1)))

wb=openpyxl.Workbook(); ws=wb.active; ws.title='Duplicate groups'
ws.append(['group','status','size_MB','copies','file_name','folder_path','kept_file_path','file_id'])
green=PatternFill('solid',fgColor='C6EFCE'); bold=Font(bold=True)
kept_count=tr_count=0
for gi,(md5,fs) in enumerate(groups,1):
    kept=[f for f in fs if f['id'] not in trashed]
    keep_path=kept[0]['path'] if kept else '(NONE!)'
    for f in sorted(fs, key=lambda x:(x['id'] not in trashed and 0 or 1)):
        is_kept = f['id'] not in trashed
        if is_kept: kept_count+=1
        else: tr_count+=1
        row=[gi,'KEPT' if is_kept else 'trashed',round(f['size']/1e6,1),len(fs),
             f['name'][:60],f['path'],keep_path if is_kept else '',f['id']]
        ws.append(row)
        if is_kept:
            for c in range(1,9): ws.cell(ws.max_row,c).fill=green; ws.cell(ws.max_row,c).font=bold
for col,w in {'A':7,'B':9,'C':10,'D':7,'E':52,'F':58,'G':58,'H':30}.items(): ws.column_dimensions[col].width=w
ws.freeze_panes='A2'

sm=wb.create_sheet('Summary')
for r in [['Metric','Value'],
    ['Duplicate groups',len(groups)],
    ['Copies kept (one per group)',kept_count],
    ['Copies trashed',tr_count],
    ['Space reclaimed (GB)',round(sum(f['size'] for f in files if f['id'] in trashed)/1e9,1)],
    ['Note','KEPT rows are green/bold; kept_file_path shows the surviving copy']]:
    sm.append(r)
sm.column_dimensions['A'].width=32; sm.column_dimensions['B'].width=60
wb.save('duplicates.xlsx')
print(f"Groups: {len(groups)} | kept: {kept_count} | trashed: {tr_count}")
