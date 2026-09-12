"""Find true duplicate files in the zBackup tree by MD5. Writes duplicates.xlsx."""
import sys, json
from collections import defaultdict
import openpyxl

files=json.load(open('drive_index_md5.json'))
by_md5=defaultdict(list)
for f in files:
    if f['md5']:
        by_md5[f['md5']].append(f)

# duplicate groups = same md5 with 2+ files
groups=[(md5,fs) for md5,fs in by_md5.items() if len(fs)>1]
# sort by reclaimable space (size * (count-1)) desc
groups.sort(key=lambda g:-(g[1][0]['size']*(len(g[1])-1)))

dup_files=sum(len(fs) for _,fs in groups)
extra_copies=sum(len(fs)-1 for _,fs in groups)
wasted=sum(fs[0]['size']*(len(fs)-1) for _,fs in groups)

wb=openpyxl.Workbook()
ws=wb.active; ws.title='Duplicate groups'
ws.append(['group','copies','size_MB','md5','file_name','folder_path','file_id','reclaimable_MB_if_dedup'])
for gi,(md5,fs) in enumerate(groups,1):
    recl=round(fs[0]['size']*(len(fs)-1)/1e6,1)
    for j,f in enumerate(sorted(fs,key=lambda x:x['path'])):
        ws.append([gi,len(fs),round(f['size']/1e6,1),md5,f['name'][:80],f['path'],f['id'],
                   recl if j==0 else ''])
for col,w in {'A':7,'B':7,'C':10,'D':34,'E':52,'F':60,'G':30,'H':12}.items():
    ws.column_dimensions[col].width=w
ws.freeze_panes='A2'

# same-folder duplicates (most clearly accidental) as a focused tab
ws2=wb.create_sheet('Same-folder dups')
ws2.append(['size_MB','copies_in_folder','file_name','folder_path','md5'])
sf=0
for md5,fs in groups:
    byfolder=defaultdict(list)
    for f in fs: byfolder[f['path']].append(f)
    for path,ff in byfolder.items():
        if len(ff)>1:
            sf+=1
            ws2.append([round(ff[0]['size']/1e6,1),len(ff),ff[0]['name'][:80],path,md5])
for col,w in {'A':10,'B':16,'C':52,'D':64,'E':34}.items(): ws2.column_dimensions[col].width=w
ws2.freeze_panes='A2'

sm=wb.create_sheet('Summary')
for r in [['Metric','Value'],
    ['Total files scanned',len(files)],
    ['Duplicate groups (same MD5, 2+ copies)',len(groups)],
    ['Files that are duplicates',dup_files],
    ['Redundant extra copies',extra_copies],
    ['Reclaimable if deduped (GB)',round(wasted/1e9,1)],
    ['Same-folder duplicate cases',sf]]:
    sm.append(r)
sm.column_dimensions['A'].width=40; sm.column_dimensions['B'].width=18
wb.save('duplicates.xlsx')
print(f"Duplicate groups: {len(groups)}")
print(f"Extra (redundant) copies: {extra_copies}")
print(f"Reclaimable if deduped: {wasted/1e9:.1f} GB")
print(f"Same-folder duplicate cases: {sf}")
